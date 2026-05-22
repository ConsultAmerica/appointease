"""
Build a short, professional team deck with app screenshots.
  python scripts/capture-presentation-screenshots.py
  python scripts/generate-team-presentation.py
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

_SCRIPT_DIR = Path(__file__).resolve().parent
if str(_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPT_DIR))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from presentation_config import LIVE_URL, OUT_DECK, SCREENSHOTS

SLATE = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)
TEAL = RGBColor(15, 118, 110)
LINE = RGBColor(226, 232, 240)
WHITE = RGBColor(255, 255, 255)

W, H = Inches(10), Inches(7.5)


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rule(slide, top: float) -> None:
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(top), Inches(8.7), Inches(0.015))
    r.fill.solid()
    r.fill.fore_color.rgb = LINE
    r.line.fill.background()


def _heading(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.5), Inches(8.7), Inches(1.0))
    tf = box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.name = "Calibri"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = SLATE
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = MUTED
        p2.space_before = Pt(4)
    _rule(slide, 1.35)


def _bullets(slide, items: list[str], top=1.55, size=17) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(top), Inches(8.7), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = line
        p.font.name = "Calibri"
        p.font.size = Pt(size)
        p.font.color.rgb = SLATE
        p.space_after = Pt(8)
        p.level = 0


def _image(slide, path: Path, left, top, width) -> bool:
    if not path.is_file():
        return False
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    return True


def slide_title(prs, notes: str) -> None:
    s = _blank(prs)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.08))
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()

    t = s.shapes.add_textbox(Inches(0.65), Inches(2.0), Inches(8.5), Inches(1.0))
    t.text_frame.text = "AppointmentAI"
    t.text_frame.paragraphs[0].font.name = "Calibri"
    t.text_frame.paragraphs[0].font.size = Pt(40)
    t.text_frame.paragraphs[0].font.bold = True
    t.text_frame.paragraphs[0].font.color.rgb = SLATE

    sub = s.shapes.add_textbox(Inches(0.65), Inches(3.0), Inches(8.5), Inches(0.6))
    sub.text_frame.text = "Multi-clinic appointment scheduling"
    sub.text_frame.paragraphs[0].font.name = "Calibri"
    sub.text_frame.paragraphs[0].font.size = Pt(18)
    sub.text_frame.paragraphs[0].font.color.rgb = MUTED

    url = s.shapes.add_textbox(Inches(0.65), Inches(4.0), Inches(8.5), Inches(0.45))
    url_tf = url.text_frame
    url_tf.text = LIVE_URL
    p = url_tf.paragraphs[0]
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.color.rgb = TEAL
    p.font.underline = True
    try:
        p.hyperlink.address = LIVE_URL
    except Exception:
        pass

    meta = s.shapes.add_textbox(Inches(0.65), Inches(6.5), Inches(8.5), Inches(0.35))
    meta.text_frame.text = "Team briefing"
    meta.text_frame.paragraphs[0].font.size = Pt(11)
    meta.text_frame.paragraphs[0].font.color.rgb = MUTED
    _notes(s, notes)


def slide_overview(prs, notes: str) -> None:
    s = _blank(prs)
    _heading(s, "Overview", "What the platform delivers")
    _bullets(
        s,
        [
            "Customers book services online with live availability.",
            "Clinic admins manage schedule, staff, and services in one dashboard.",
            "Optional AI chat uses the same booking data as the web flows.",
        ],
    )
    _notes(s, notes)


def slide_image(prs, title: str, subtitle: str, img: Path, notes: str) -> None:
    s = _blank(prs)
    _heading(s, title, subtitle)
    if _image(s, img, 0.65, 1.5, 8.7):
        pass
    else:
        ph = s.shapes.add_textbox(Inches(0.65), Inches(3.2), Inches(8.7), Inches(1))
        ph.text_frame.text = f"[Screenshot missing: {img.name}]"
        ph.text_frame.paragraphs[0].font.color.rgb = MUTED
    _notes(s, notes)


def slide_roles(prs, notes: str) -> None:
    s = _blank(prs)
    _heading(s, "Users", "Three roles, separate workspaces")
    cols = [
        ("Customer", "Book visits, view and cancel appointments."),
        ("Staff", "See assigned visits and availability."),
        ("Admin", "Run the clinic: services, staff, today’s schedule."),
    ]
    x = 0.65
    for title, body in cols:
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(2.85), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = LINE
        tb = s.shapes.add_textbox(Inches(x + 0.2), Inches(1.85), Inches(2.45), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.text = title
        tf.paragraphs[0].font.name = "Calibri"
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = TEAL
        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.name = "Calibri"
        p2.font.size = Pt(13)
        p2.font.color.rgb = SLATE
        p2.space_before = Pt(10)
        x += 3.05
    _notes(s, notes)


def slide_stack(prs, notes: str) -> None:
    s = _blank(prs)
    _heading(s, "Technical foundation", None)
    rows = [
        ("Application", "Next.js 16, React, TypeScript, Tailwind"),
        ("Data", "PostgreSQL (Neon), Prisma ORM"),
        ("Auth & security", "Auth.js, role-based access, audit logging"),
        ("Hosting", "Vercel serverless"),
    ]
    y = 1.7
    for label, val in rows:
        lb = s.shapes.add_textbox(Inches(0.65), Inches(y), Inches(2.2), Inches(0.4))
        lb.text_frame.text = label
        lb.text_frame.paragraphs[0].font.name = "Calibri"
        lb.text_frame.paragraphs[0].font.size = Pt(14)
        lb.text_frame.paragraphs[0].font.bold = True
        lb.text_frame.paragraphs[0].font.color.rgb = TEAL
        vb = s.shapes.add_textbox(Inches(2.9), Inches(y), Inches(6.4), Inches(0.5))
        vb.text_frame.text = val
        vb.text_frame.paragraphs[0].font.name = "Calibri"
        vb.text_frame.paragraphs[0].font.size = Pt(14)
        vb.text_frame.paragraphs[0].font.color.rgb = SLATE
        y += 0.75
    _notes(s, notes)


def slide_close(prs, notes: str) -> None:
    s = _blank(prs)
    t = s.shapes.add_textbox(Inches(0.65), Inches(2.8), Inches(8.7), Inches(0.8))
    t.text_frame.text = "Questions?"
    t.text_frame.paragraphs[0].font.name = "Calibri"
    t.text_frame.paragraphs[0].font.size = Pt(36)
    t.text_frame.paragraphs[0].font.bold = True
    t.text_frame.paragraphs[0].font.color.rgb = SLATE
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    u = s.shapes.add_textbox(Inches(0.65), Inches(4.0), Inches(8.7), Inches(0.5))
    u.text_frame.text = LIVE_URL
    u.text_frame.paragraphs[0].font.name = "Calibri"
    u.text_frame.paragraphs[0].font.size = Pt(14)
    u.text_frame.paragraphs[0].font.color.rgb = TEAL
    u.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    try:
        u.text_frame.paragraphs[0].hyperlink.address = LIVE_URL
    except Exception:
        pass

    d = s.shapes.add_textbox(Inches(0.65), Inches(5.0), Inches(8.7), Inches(0.8))
    d.text_frame.text = "Demo: admin@demo-clinic.com  ·  Demo Wellness Clinic in seed data"
    d.text_frame.paragraphs[0].font.name = "Calibri"
    d.text_frame.paragraphs[0].font.size = Pt(12)
    d.text_frame.paragraphs[0].font.color.rgb = MUTED
    d.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    _notes(s, notes)


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    shots = SCREENSHOTS

    slide_title(
        prs,
        "Introduce AppointmentAI as a production scheduling product. Open the live URL if available.",
    )
    slide_overview(
        prs,
        "Keep this slide under one minute. Focus on outcomes, not feature lists.",
    )
    slide_image(
        prs,
        "Public site",
        "Landing and entry to booking",
        shots / "01-home.png",
        "Walk through hero and primary call-to-action.",
    )
    slide_image(
        prs,
        "Booking",
        "Clinic, service, and time selection",
        shots / "02-book.png",
        "Show live slots loaded from the database.",
    )
    slide_image(
        prs,
        "Admin dashboard",
        "Day-to-day operations",
        shots / "03-admin.png",
        "Log in as clinic admin. Cover today’s list and key actions.",
    )
    slide_roles(
        prs,
        "Clarify who signs in where: /customer, /staff, /admin.",
    )
    slide_stack(
        prs,
        "For engineering stakeholders only; skip or shorten for business audiences.",
    )
    slide_close(
        prs,
        "Leave time for questions. Offer a short live demo if the environment is up.",
    )

    return prs


def main():
    OUT_DECK.parent.mkdir(parents=True, exist_ok=True)
    if len(list(SCREENSHOTS.glob("*.png"))) < 2:
        print("Capturing screenshots (start npm run dev if using localhost)…")
        subprocess.run([sys.executable, str(_SCRIPT_DIR / "capture-presentation-screenshots.py")], check=False)

    prs = build()
    try:
        prs.save(str(OUT_DECK))
        print(f"Created: {OUT_DECK}")
    except PermissionError:
        alt = OUT_DECK.with_stem(OUT_DECK.stem + "-new")
        prs.save(str(alt))
        print(f"Original file locked. Created: {alt}")


if __name__ == "__main__":
    main()
